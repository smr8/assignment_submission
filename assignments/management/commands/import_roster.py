"""
将花名册数据批量导入系统

用法:
  python manage.py import_roster                                   # 使用内置名单
  python manage.py import_roster --file 班级名单.et                 # 从文件导入
  python manage.py import_roster --file 名单.xlsx --dry-run         # 只预览不写库
  python manage.py import_roster --file 名单.csv --update-names     # 同时用表里的姓名覆盖已有用户

支持的名单格式: .csv / .tsv / .xls / .et(WPS 表格) / .xlsx
表头需要包含「学号」「姓名」两列（列顺序不限，缺表头时按内容自动识别）。
"""
import csv
import os
import re

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from accounts.models import User

DEFAULT_PASSWORD = "123456"

# 未指定文件时使用的内置名单
ROSTER = [
    ("2023210834", "严菲"),
    ("2023210857", "王政"),
    ("2024210834", "刘坤弘"),
    ("2024210835", "曾子耕"),
    ("2024210836", "刘砚泽"),
    ("2024210837", "党囡囡"),
    ("2024210838", "李雨泽"),
    ("2024210839", "商静怡"),
    ("2024210840", "张翼鹏"),
    ("2024210841", "于文杰"),
    ("2024210842", "杨锦雯"),
    ("2024210844", "李心妍"),
    ("2024210845", "高瑞雪"),
    ("2024210846", "张靓靓"),
    ("2024210847", "孙荣昱"),
    ("2024210848", "袁梦豪"),
    ("2024210849", "常真"),
    ("2024210850", "朱仁浩"),
    ("2024210851", "舒子达"),
    ("2024210852", "易靖琦"),
    ("2024210853", "姜冠宇"),
    ("2024210855", "舒明睿"),
    ("2024210856", "罗艾琳"),
    ("2024210857", "措姆"),
    ("2024210859", "雷雨声"),
    ("2024210860", "周崇一"),
    ("2024210862", "陈晨"),
    ("2024210863", "李黄灏"),
    ("2024210864", "刘铭"),
    ("2024210865", "朱泓锦"),
    ("2024210866", "陈建"),
    ("2024210868", "刘贤翔"),
    ("2024210869", "江程程"),
    ("2024210870", "杨一林"),
    ("2024210871", "佟佳乐"),
    ("2024210872", "陈宇扬"),
    ("2024210873", "姚佳旭"),
    ("2024210876", "宋佳鑫"),
    ("2024210949", "盛宇航"),
    ("2024211050", "徐浩字"),
    ("2024211069", "邱畅"),
    ("2024212443", "孙嘉慧"),
    ("2024212444", "韦吉月"),
    ("2024212447", "张伊宁"),
    ("2024213595", "吴佩谣"),
    ("2024214094", "田硕"),
    ("2024214348", "于杭左"),
    ("2024214380", "李云燕"),
]

ID_HEADERS = {"学号", "学籍号", "学生学号", "编号", "id", "student_id", "studentid"}
NAME_HEADERS = {"姓名", "名字", "学生姓名", "name"}
CJK_RE = re.compile(r"^[\u4e00-\u9fa5·]{2,8}$")
ID_RE = re.compile(r"^\d{6,}$")


def normalize_id(value):
    """把单元格里的学号统一成字符串。

    Excel/WPS 里学号常被存成数字，xlrd 会读成 float（2411001332.0），
    直接 str() 会得到 '2411001332.0'，必须单独处理。
    """
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    text = str(value).strip().replace(" ", "").replace("\u3000", "")
    # 文本形式但带小数点的情况，例如 "2411001332.0"
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".")[0]
    return text


class Command(BaseCommand):
    help = "导入花名册学生信息（支持 csv/tsv/xls/et/xlsx）"

    def add_arguments(self, parser):
        parser.add_argument("--file", dest="file", default="", help="名单文件路径")
        parser.add_argument("--password", dest="password", default=DEFAULT_PASSWORD,
                            help=f"新用户初始密码（默认 {DEFAULT_PASSWORD}）")
        parser.add_argument("--update-names", action="store_true",
                            help="已存在的用户也按表里的姓名更新（默认跳过）")
        parser.add_argument("--dry-run", action="store_true", help="只解析并预览，不写入数据库")

    def handle(self, *args, **options):
        file_path = options["file"]
        password = options["password"]
        update_names = options["update_names"]
        dry_run = options["dry_run"]

        if file_path:
            if not os.path.exists(file_path):
                raise CommandError(f"文件不存在: {file_path}")
            rows = self._read_rows(file_path)
            pairs, header_row = extract_pairs(rows)
            self.stdout.write(f"📄 文件: {file_path}")
            if header_row is not None:
                self.stdout.write(f"   解析到 {len(rows)} 行，表头在第 {header_row} 行")
            else:
                self.stdout.write(f"   解析到 {len(rows)} 行（未发现表头，按内容识别列）")
        else:
            pairs = [(normalize_id(sid), name.strip()) for sid, name in ROSTER]
            self.stdout.write("📄 使用内置名单")

        if not pairs:
            raise CommandError("没有解析到任何「学号 + 姓名」记录，请检查文件格式。")

        duplicates = self._find_duplicates(pairs)
        if duplicates:
            self.stdout.write(self.style.WARNING(f"⚠️  名单内重复学号: {duplicates}"))

        self.stdout.write(f"   有效记录: {len(pairs)} 条")
        if dry_run:
            self.stdout.write("\n[预览模式] 将导入以下记录：")
            for sid, name in pairs:
                exists = User.objects.filter(student_id=sid).exists()
                mark = "已存在→跳过" if exists else "新建"
                self.stdout.write(f"  {sid}  {name}   [{mark}]")
            self.stdout.write(self.style.WARNING("\n未写入数据库（--dry-run）。"))
            return

        created = skipped = renamed = 0
        with transaction.atomic():
            for student_id, name in pairs:
                user = User.objects.filter(student_id=student_id).first()
                if user is None:
                    User.objects.create_user(
                        student_id=student_id, name=name, password=password,
                        is_staff=False, is_active=True,
                    )
                    self.stdout.write(self.style.SUCCESS(f"已创建: {student_id} {name}"))
                    created += 1
                    continue

                if update_names and user.name != name:
                    old = user.name
                    user.name = name
                    user.save(update_fields=["name"])
                    self.stdout.write(self.style.WARNING(f"已改名: {student_id} {old} → {name}"))
                    renamed += 1
                else:
                    self.stdout.write(self.style.WARNING(f"跳过已存在: {student_id} {name}"))
                    skipped += 1

        self.stdout.write("")
        self.stdout.write(self.style.SUCCESS(
            f"导入完成: 新增 {created} 人, 改名 {renamed} 人, 跳过 {skipped} 人"
        ))
        self.stdout.write(f"学生总数: {User.objects.filter(is_staff=False).count()} 人")

    # ---------------- 解析 ----------------

    def _read_rows(self, path):
        ext = os.path.splitext(path)[1].lower()
        if ext in (".csv", ".tsv", ".txt"):
            return self._read_delimited(path, ext)
        if ext in (".xls", ".et", ".ett"):
            return self._read_xls(path)
        if ext in (".xlsx", ".xlsm"):
            return self._read_xlsx(path)
        raise CommandError(f"不支持的文件类型: {ext}（支持 csv/tsv/xls/et/xlsx）")

    def _read_delimited(self, path, ext):
        for encoding in ("utf-8-sig", "gbk"):
            try:
                with open(path, "r", encoding=encoding, newline="") as fh:
                    sample = fh.read(4096)
                    fh.seek(0)
                    delimiter = "\t" if ext == ".tsv" else (
                        csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
                        if sample.strip() else ","
                    )
                    return [[c.strip() for c in row] for row in csv.reader(fh, delimiter=delimiter)]
            except UnicodeDecodeError:
                continue
        raise CommandError(f"无法解码文件（尝试过 utf-8 / gbk）: {path}")

    def _read_xls(self, path):
        try:
            import xlrd
        except ImportError:
            raise CommandError("读取 .xls/.et 需要 xlrd，请先执行: pip install xlrd")
        try:
            book = xlrd.open_workbook(path)
        except Exception as exc:
            raise CommandError(f"无法解析表格（WPS 的 .et 请另存为 .xlsx 或用 xlrd 读取）: {exc}")
        sheet = book.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]

    def _read_xlsx(self, path):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise CommandError("读取 .xlsx 需要 openpyxl，请先执行: pip install openpyxl")
        book = load_workbook(path, read_only=True, data_only=True)
        sheet = book.active
        return [[c for c in row] for row in sheet.iter_rows(values_only=True)]

    def _find_duplicates(self, pairs):
        seen, dups = set(), []
        for sid, _ in pairs:
            if sid in seen:
                dups.append(sid)
            seen.add(sid)
        return dups


def extract_pairs(rows):
    """从二维表里抽出 (学号, 姓名)，返回 (pairs, 表头行号或 None)"""
    # 1) 找表头
    header_row = id_col = name_col = None
    for index, row in enumerate(rows[:20]):
        cells = [str(c).strip() for c in row]
        found_id = next((i for i, c in enumerate(cells) if c.lower() in ID_HEADERS), None)
        found_name = next((i for i, c in enumerate(cells) if c.lower() in NAME_HEADERS), None)
        if found_id is not None and found_name is not None:
            header_row, id_col, name_col = index, found_id, found_name
            break

    # 2) 没表头就按内容识别列
    if header_row is None:
        id_col, name_col = _guess_columns(rows)
        data_rows = rows
    else:
        data_rows = rows[header_row + 1:]

    pairs = []
    for row in data_rows:
        if id_col >= len(row) or name_col >= len(row):
            continue
        student_id = normalize_id(row[id_col])
        name = str(row[name_col]).strip().replace("\u3000", "")
        if not student_id or not name:
            continue
        # 过滤掉「合计」之类的汇总行
        if not ID_RE.match(student_id):
            continue
        if name.lower() in NAME_HEADERS:
            continue
        pairs.append((student_id, name))

    return pairs, header_row


def _guess_columns(rows):
    """无表头时：学号列 = 数字型且长度>=6 最多的列；姓名列 = 中文短词最多的列"""
    width = max((len(r) for r in rows), default=0)
    id_scores, name_scores = {}, {}
    for col in range(width):
        id_scores[col] = name_scores[col] = 0
        for row in rows:
            if col >= len(row):
                continue
            value = normalize_id(row[col])
            if ID_RE.match(value):
                id_scores[col] += 1
            if CJK_RE.match(str(row[col]).strip()):
                name_scores[col] += 1
    id_col = max(id_scores, key=id_scores.get) if any(id_scores.values()) else 0
    name_candidates = {c: s for c, s in name_scores.items() if c != id_col and s > 0}
    name_col = max(name_candidates, key=name_candidates.get) if name_candidates else id_col + 1
    return id_col, name_col
